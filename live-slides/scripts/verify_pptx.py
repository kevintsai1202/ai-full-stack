"""
驗證產出的 PPTX 是否符合交付條件。

檢查項目：
1. 頁數與 index.html 的 .slide 數量一致
2. 每頁都有講者備註（對應 data-note）
3. 文字是原生可編輯的（有實際的文字框，不是整頁圖片）
4. 所有截圖都嵌進檔案裡
5. 備援影片以 MP4 嵌入（PowerPoint 不支援 WebM）

執行：
    python live-slides/scripts/verify_pptx.py

驗證其他 deck（參數與 html_to_pptx.py 一致）：
    python live-slides/scripts/verify_pptx.py --html course-package/slides/00-course-orientation.html
"""
from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

# 預設值維持原直播簡報，讓既有的無參數呼叫行為完全不變
SLIDES_DIR = Path(__file__).resolve().parent.parent
HTML_PATH = SLIDES_DIR / "index.html"
PPTX_PATH = SLIDES_DIR / "別再問AI能不能做.pptx"

failures: list[str] = []


def check(name: str, passed: bool, detail: str = "") -> None:
    """記錄一項檢查結果。"""
    print(f"{'  OK  ' if passed else ' FAIL '} {name}{f' — {detail}' if detail else ''}")
    if not passed:
        failures.append(name)


def _parse_args():
    """解析命令列參數；兩個都省略時等同原本的無參數行為。"""
    parser = argparse.ArgumentParser(description="驗證 PPTX 是否符合交付條件")
    parser.add_argument("--html", help=f"原稿 HTML deck（預設：{HTML_PATH}）")
    parser.add_argument("--pptx", help="待驗證的 PPTX（預設：與 HTML 同名的 .pptx）")
    return parser.parse_args()


def main() -> int:
    global SLIDES_DIR, HTML_PATH, PPTX_PATH
    args = _parse_args()
    if args.html:
        HTML_PATH = Path(args.html).resolve()
        SLIDES_DIR = HTML_PATH.parent
        PPTX_PATH = HTML_PATH.with_suffix(".pptx")
    if args.pptx:
        PPTX_PATH = Path(args.pptx).resolve()

    html = HTML_PATH.read_text(encoding="utf-8")
    expected = len(re.findall(r'<section class="slide', html))
    prs = Presentation(PPTX_PATH)
    slides = prs.slides

    check("頁數與原稿一致", len(slides._sldIdLst) == expected,
          f"PPTX {len(slides._sldIdLst)} 頁 / HTML {expected} 頁")

    # 講者備註：每頁都要有，否則直播時該頁沒有提示
    missing_notes = [i for i, s in enumerate(slides, 1)
                     if not (s.has_notes_slide
                             and s.notes_slide.notes_text_frame.text.strip())]
    check("每頁都有講者備註", not missing_notes,
          f"缺第 {missing_notes} 頁" if missing_notes else "")

    # 文字可編輯：每頁至少要有一個含文字的文字框
    no_text = []
    total_runs = 0
    for index, slide in enumerate(slides, 1):
        runs = [run for shape in slide.shapes if shape.has_text_frame
                for para in shape.text_frame.paragraphs for run in para.runs
                if run.text.strip()]
        total_runs += len(runs)
        if not runs:
            no_text.append(index)
    check("每頁都有可編輯文字", not no_text,
          f"共 {total_runs} 個文字 run" if not no_text else f"第 {no_text} 頁無文字")

    # 媒體檔：截圖與影片都要真的打包進 pptx
    with zipfile.ZipFile(PPTX_PATH) as zf:
        media = [n for n in zf.namelist() if n.startswith("ppt/media/")]
    images = [m for m in media if m.lower().endswith((".png", ".jpg", ".jpeg"))]
    videos = [m for m in media if m.lower().endswith(".mp4")]
    html_imgs = len(set(re.findall(r'<img src="(assets/[^"]+)"', html)))

    check("截圖已嵌入", len(images) >= html_imgs,
          f"pptx {len(images)} 張 / html 引用 {html_imgs} 張")
    # 只有原稿真的放了 <video> 才要求 MP4：無影片的 deck（例如課程包開場）
    # 若照樣要求，會回報一個「缺了本來就不存在的東西」的假失敗
    if "<video" in html:
        check("備援影片以 MP4 嵌入", len(videos) >= 1, ", ".join(videos))
    else:
        print("  --   備援影片以 MP4 嵌入 — 原稿無影片，略過")
    check("未殘留 WebM", not [m for m in media if m.endswith(".webm")])

    # 版面：所有圖形都要落在頁面內，內容太高被推出底部等於觀眾看不到
    # 底部列固定貼齊頁底，不在檢查之列
    slide_w, slide_h = prs.slide_width, prs.slide_height
    bar_top = slide_h - int(slide_h * 46 / 900)
    overflow = []
    for index, slide in enumerate(slides, 1):
        for shape in slide.shapes:
            if shape.top is None or shape.height is None:
                continue
            if shape.top >= bar_top:
                continue
            if shape.top < 0 or shape.top + shape.height > bar_top + int(slide_h * 4 / 900):
                overflow.append(index)
                break
    check("所有圖形都在頁面內", not overflow,
          f"第 {overflow} 頁有圖形超出底部或頂端" if overflow else "")

    # 字級：投影用簡報，內文不得低於 12pt（＝原稿 20px），底部列除外
    tiny = {}
    for index, slide in enumerate(slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame or (shape.top is not None and shape.top >= bar_top):
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip() and run.font.size is not None and run.font.size.pt < 12:
                        tiny.setdefault(index, set()).add(run.font.size.pt)
    check("內文字級不低於 12pt", not tiny,
          "；".join(f"第 {p} 頁 {sorted(s)}pt" for p, s in sorted(tiny.items())) if tiny else "")

    # 原稿有 QR Code 的頁面（待機頁、掃碼加入頁），PPTX 該頁一定要有圖片
    qr_pages = [i for i, m in enumerate(re.finditer(r'<section class="slide[^>]*>(.*?)</section>', html, re.S), 1)
                if "exam-join-qr.png" in m.group(1)]
    missing_qr = [p for p in qr_pages
                  if not any(shape.shape_type == 13 for shape in list(slides)[p - 1].shapes)]
    check("QR Code 頁含圖片", not missing_qr,
          f"原稿第 {qr_pages} 頁有 QR，PPTX 第 {missing_qr} 頁缺圖" if missing_qr else f"第 {qr_pages} 頁")

    print()
    if failures:
        print(f"驗證失敗 {len(failures)} 項：{'、'.join(failures)}", file=sys.stderr)
        return 1
    print("PPTX 驗證全部通過。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
